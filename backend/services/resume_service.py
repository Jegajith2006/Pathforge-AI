"""
PathForge AI - Resume Parsing & Evidence Extraction Service
Extracts text from PDF, DOCX, and TXT files, normalizes formatting, executes
Gemini structured entity extraction with neural BGE fallback, maps skills to database
ontology, and integrates exposure evidence into the centralized mastery system.

CORE PRINCIPLE:
Resume detection indicates 'evidence of exposure' (Developing baseline score ~45-65),
NOT 'verified mastery'. Final mastery is computed via centralized multi-source calibration
and cannot achieve 'Advanced' status solely from resume mentions.
"""

import os
import re
import uuid
import json
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple
from sqlalchemy.orm import Session

from config import settings
from database.models import Student, Skill, ResumeAnalysis, StudentSkill, ResumeDocument
from ml.model_config import RESUME_SIMILARITY_THRESHOLD
from ml.skill_matching import (
    semantic_skill_detection,
    calculate_final_mastery,
)
from services.skill_gap_service import update_or_create_student_skill
from services.gemini_service import gemini_service
from schemas.resume import (
    ResumeAnalysisResponse,
    ResumeExtractedSkill,
    ResumeExtractedProject,
    ResumeExtractedCertificate,
    ResumeExtractedExperience,
    ResumeExtractedEducation,
)

logger = logging.getLogger("pathforge.resume")


def normalize_text(raw_text: str) -> str:
    """
    Normalizes resume text:
      - Normalizes whitespace and line breaks
      - Cleans up non-printable control characters
      - Preserves structural section boundaries
      - Safely limits length to 15,000 characters
    """
    if not raw_text:
        return ""

    # Replace null bytes and non-printable control characters
    cleaned = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]", " ", raw_text)
    
    # Normalize multiple horizontal whitespaces
    cleaned = re.sub(r"[ \t]+", " ", cleaned)
    
    # Normalize multiple newlines (max 2 consecutive newlines)
    cleaned = re.sub(r"\n\s*\n\s*\n+", "\n\n", cleaned)
    
    # Clean leading/trailing whitespaces per line
    lines = [line.strip() for line in cleaned.split("\n")]
    normalized = "\n".join(lines).strip()

    # Controlled limit to avoid overwhelming downstream context
    return normalized[:15000]


def extract_text_from_pdf(file_path: Path) -> str:
    """
    Extracts text from PDF using pypdf.
    Preserves document structure and detects scanned/image-only PDFs.
    """
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("pypdf is required for PDF extraction.")

    try:
        reader = PdfReader(str(file_path))
        page_count = len(reader.pages)
        if page_count == 0:
            raise ValueError("Uploaded PDF document contains 0 pages.")

        extracted_pages = []
        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                extracted_pages.append(page_text.strip())

        combined_text = "\n\n".join(extracted_pages).strip()

        # Scanned/Image-only PDF detection:
        if page_count > 0 and len(combined_text) < 15:
            raise ValueError("Scanned/image-only PDF requires OCR")

        return combined_text
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Corrupted or unreadable PDF file: {str(e)}")


def extract_text_from_docx(file_path: Path) -> str:
    """
    Extracts text from DOCX using python-docx.
    Extracts paragraphs, headings, bullet points, and tables.
    """
    try:
        import docx
    except ImportError:
        raise RuntimeError("python-docx is required for DOCX extraction.")

    try:
        doc = docx.Document(str(file_path))
        elements = []

        # Extract text from paragraphs and headings
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                elements.append(text)

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    elements.append(" | ".join(row_cells))

        combined_text = "\n".join(elements).strip()
        if not combined_text:
            raise ValueError("Uploaded DOCX document contains no readable text content.")

        return combined_text
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Corrupted or unreadable DOCX file: {str(e)}")


def extract_text_from_file(file_path: Path) -> str:
    """
    Extracts text from supported file types (PDF, DOCX, TXT, MD) with validation.
    """
    if not file_path.exists():
        raise ValueError(f"File not found: {file_path}")

    filesize = file_path.stat().st_size
    if filesize == 0:
        raise ValueError("Uploaded file is empty (0 bytes).")

    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        raw_text = extract_text_from_pdf(file_path)
    elif suffix == ".docx":
        raw_text = extract_text_from_docx(file_path)
    elif suffix in [".txt", ".md", ".csv"]:
        try:
            raw_text = file_path.read_text(encoding="utf-8", errors="ignore")
            if not raw_text.strip():
                raise ValueError("Uploaded text file is empty.")
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"Unable to read text file: {str(e)}")
    else:
        raise ValueError(f"Unsupported file format '{suffix}'. Allowed: PDF, DOCX, TXT.")

    return normalize_text(raw_text)


# Known alias map for skill normalization
SKILL_ALIASES = {
    "scikit learn": "Scikit-learn",
    "scikit-learn": "Scikit-learn",
    "sklearn": "Scikit-learn",
    "ml": "Machine Learning",
    "machine learning": "Machine Learning",
    "deep learning": "Deep Learning",
    "dl": "Deep Learning",
    "nlp": "Natural Language Processing",
    "cv": "Computer Vision",
    "pytorch framework": "PyTorch",
    "pytorch": "PyTorch",
    "torch": "PyTorch",
    "tensorflow": "TensorFlow",
    "tf": "TensorFlow",
    "postgres": "PostgreSQL",
    "postgresql": "PostgreSQL",
    "psql": "PostgreSQL",
    "js": "JavaScript",
    "ts": "TypeScript",
    "docker container": "Docker",
    "docker": "Docker",
    "k8s": "Kubernetes",
    "kubernetes": "Kubernetes",
    "rest api": "REST APIs",
    "fastapi": "FastAPI",
    "flask": "Flask",
    "django": "Django",
    "react": "React",
    "react.js": "React",
    "reactjs": "React",
    "next.js": "Next.js",
    "nextjs": "Next.js",
    "node": "Node.js",
    "nodejs": "Node.js",
    "aws cloud": "AWS",
    "amazon web services": "AWS",
    "gcp cloud": "GCP",
    "google cloud": "GCP",
    "google cloud platform": "GCP",
    "azure cloud": "Azure",
    "microsoft azure": "Azure",
    "spark": "Apache Spark",
    "apache spark": "Apache Spark",
}


def normalize_extracted_skills(
    db: Session,
    raw_skills: List[str],
    resume_text: str,
) -> Tuple[List[Dict[str, Any]], List[str]]:
    """
    Normalizes Gemini extracted skill names against PathForge database skills table.
    Uses exact matching, alias lookup, and BGE dense embedding similarity.
    """
    all_db_skills = db.query(Skill).all()
    db_skill_map = {s.skill_name.lower(): s for s in all_db_skills}
    
    candidate_skills = [
        {
            "skill_id": s.skill_id,
            "skill_name": s.skill_name,
            "category": s.category,
            "description": s.description,
        }
        for s in all_db_skills
    ]

    matched_dict: Dict[str, Dict[str, Any]] = {}
    unmatched_skills: List[str] = []

    # 1. Match Gemini raw skills against DB ontology
    for raw in raw_skills:
        raw_clean = raw.strip()
        if not raw_clean:
            continue

        raw_lower = raw_clean.lower()
        alias_target = SKILL_ALIASES.get(raw_lower, raw_clean)
        alias_lower = alias_target.lower()

        if alias_lower in db_skill_map:
            s_obj = db_skill_map[alias_lower]
            matched_dict[s_obj.skill_id] = {
                "skill_id": s_obj.skill_id,
                "skill_name": s_obj.skill_name,
                "confidence": 0.95,
                "evidence_type": "Gemini Structured Extraction",
            }
        elif raw_lower in db_skill_map:
            s_obj = db_skill_map[raw_lower]
            matched_dict[s_obj.skill_id] = {
                "skill_id": s_obj.skill_id,
                "skill_name": s_obj.skill_name,
                "confidence": 0.90,
                "evidence_type": "Gemini Structured Extraction",
            }
        else:
            unmatched_skills.append(raw_clean)

    # 2. Execute BGE semantic + regex detection on resume text for high-fidelity ontology coverage
    semantic_matches = semantic_skill_detection(
        resume_text=resume_text,
        candidate_skills=candidate_skills,
        threshold=RESUME_SIMILARITY_THRESHOLD,
    )

    for sm in semantic_matches:
        s_id = sm["skill_id"]
        conf = float(sm.get("confidence", 0.88))
        if s_id not in matched_dict:
            matched_dict[s_id] = {
                "skill_id": s_id,
                "skill_name": sm["skill_name"],
                "confidence": conf,
                "evidence_type": "BGE Neural Semantic Extraction",
            }
        else:
            matched_dict[s_id]["confidence"] = max(matched_dict[s_id]["confidence"], conf)

    return list(matched_dict.values()), unmatched_skills


def parse_and_analyze_resume(
    db: Session,
    student_id: str,
    file_path: Path,
    original_filename: str,
    resume_id: Optional[str] = None,
) -> Dict[str, Any]:
    """
    Full Resume Analysis Pipeline:
      1. Validates and extracts text from PDF/DOCX/TXT
      2. Normalizes text
      3. Invokes Gemini structured extraction (with resilient fallback)
      4. Normalizes detected skills against PathForge database ontology
      5. Persists ResumeAnalysis records
      6. Calibrates student_skills.resume_score (exposure baseline ~45-65)
      7. Saves persistent ResumeDocument history record
      8. Recalculates final evidence mastery and returns structured response
    """
    student = db.query(Student).filter(Student.student_id == student_id).first()
    if not student:
        raise ValueError(f"Student '{student_id}' not found.")

    # 1. Extract & normalize text
    text = extract_text_from_file(file_path)

    # 2. Gemini Structured Extraction
    structured_data = gemini_service.extract_structured_resume_data(text)

    # 3. Skill Normalization
    matched_skills, unmatched_skills = normalize_extracted_skills(
        db=db,
        raw_skills=structured_data.get("skills", []),
        resume_text=text,
    )

    if not matched_skills:
        raise ValueError("Could not extract skills from document. Please ensure the document contains relevant technical experience, skills, or project details.")

    # Check evidence indicators in text
    project_mentions = len(structured_data.get("projects", [])) or len(
        re.findall(r"\b(project|pipeline|architecture|model|system|platform|repository)\b", text, re.IGNORECASE)
    )
    cert_mentions = len(structured_data.get("certificates", [])) or len(
        re.findall(r"\b(certificate|certified|course|specialization|credential)\b", text, re.IGNORECASE)
    )
    exp_mentions = len(structured_data.get("experience", []))

    detected_skills_list = []

    for m in matched_skills:
        s_id = m["skill_id"]
        s_name = m["skill_name"]
        conf = float(m.get("confidence", 0.90))

        # Store or update in ResumeAnalysis table
        res_record = (
            db.query(ResumeAnalysis)
            .filter(
                ResumeAnalysis.student_id == student_id,
                ResumeAnalysis.skill_id == s_id,
            )
            .first()
        )

        if not res_record:
            res_record = ResumeAnalysis(
                resume_id=f"RES-{uuid.uuid4().hex[:8].upper()}",
                student_id=student_id,
                skill_id=s_id,
                detected=True,
                confidence=conf,
                evidence_type=m.get("evidence_type", "Resume Extraction"),
                project_evidence=bool(project_mentions > 0),
                experience_evidence=bool(exp_mentions > 0),
                certificate_evidence=bool(cert_mentions > 0),
            )
            db.add(res_record)
        else:
            res_record.confidence = conf
            res_record.detected = True
            res_record.project_evidence = bool(project_mentions > 0)
            res_record.experience_evidence = bool(exp_mentions > 0)
            res_record.certificate_evidence = bool(cert_mentions > 0)

        # Resume exposure score: bounded between 45 and 65
        # (Resume evidence is only exposure, NOT verified mastery)
        resume_exposure_score = int(round(50 + conf * 15))
        resume_exposure_score = min(65, max(45, resume_exposure_score))

        # Update student_skills using central multi-source evidence calculation
        updated_ss = update_or_create_student_skill(
            db=db,
            student_id=student_id,
            skill_id=s_id,
            resume_score=resume_exposure_score,
        )

        detected_skills_list.append({
            "skill_id": s_id,
            "skill_name": s_name,
            "confidence": conf,
            "evidence_type": m.get("evidence_type", "Resume Extraction"),
            "project_evidence": bool(project_mentions > 0),
            "experience_evidence": bool(exp_mentions > 0),
            "certificate_evidence": bool(cert_mentions > 0),
            "resume_score": resume_exposure_score,
            "final_mastery": updated_ss.final_mastery,
            "mastery_level": updated_ss.mastery_level,
        })

    db.commit()

    # Format extracted projects, certificates, experience, education
    extracted_projects = [
        ResumeExtractedProject(
            title=p.get("title", "Project"),
            description=p.get("description", ""),
            technologies=p.get("technologies", []),
            duration=p.get("duration"),
        )
        for p in structured_data.get("projects", [])
    ]

    extracted_certificates = [
        ResumeExtractedCertificate(
            name=c.get("name", "Certificate"),
            issuer=c.get("issuer", ""),
            date=c.get("date"),
            related_skills=c.get("related_skills", []),
        )
        for c in structured_data.get("certificates", [])
    ]

    extracted_experience = [
        ResumeExtractedExperience(
            company=e.get("company", "Company"),
            role=e.get("role", "Engineer"),
            duration=e.get("duration"),
            description=e.get("description", ""),
            technologies=e.get("technologies", []),
        )
        for e in structured_data.get("experience", [])
    ]

    extracted_education = [
        ResumeExtractedEducation(
            institution=ed.get("institution", "University"),
            degree=ed.get("degree", "Degree"),
            field=ed.get("field"),
            year=ed.get("year"),
        )
        for ed in structured_data.get("education", [])
    ]

    # Generate comprehensive, structured resume summary strictly from extracted document content
    prof_summary = structured_data.get("summary") or "Technical candidate profile parsed from uploaded document."
    edu_str = ", ".join([f"{ed.degree} ({ed.institution})" for ed in extracted_education]) if extracted_education else "Not detected in document"
    top_skills_str = ", ".join([s["skill_name"] for s in detected_skills_list[:8]]) if detected_skills_list else "None detected"
    proj_str = ", ".join([p.title for p in extracted_projects[:3]]) if extracted_projects else "No standalone project artifacts detected"
    exp_str = ", ".join([f"{e.role} at {e.company}" for e in extracted_experience[:2]]) if extracted_experience else "No prior work experience detected"
    cert_str = ", ".join([c.name for c in extracted_certificates[:3]]) if extracted_certificates else "No certifications detected"

    student_record = db.query(Student).filter(Student.student_id == student_id).first()
    missing_career_skills_str = "Select a target career to identify career-specific missing skills."
    if student_record and student_record.target_career_id:
        from database.models import CareerSkill
        c_skills = db.query(CareerSkill).filter(CareerSkill.career_id == student_record.target_career_id).all()
        detected_names_lower = {s["skill_name"].lower() for s in detected_skills_list}
        missing_skills = [
            cs.skill.skill_name if cs.skill else cs.skill_id
            for cs in c_skills
            if (cs.skill.skill_name if cs.skill else cs.skill_id).lower() not in detected_names_lower
        ]
        if missing_skills:
            missing_career_skills_str = ", ".join(missing_skills[:5])
        else:
            missing_career_skills_str = "All target career core skills detected in document exposure."

    summary_text = (
        f"1. Professional Summary:\n{prof_summary}\n\n"
        f"2. Education:\n{edu_str}\n\n"
        f"3. Work & Internship Experience:\n{exp_str}\n\n"
        f"4. Extracted Technical Skills:\n{top_skills_str}\n\n"
        f"5. Documented Projects:\n{proj_str}\n\n"
        f"6. Detected Certifications:\n{cert_str}\n\n"
        f"7. Key Strengths:\n{len(detected_skills_list)} technical competencies detected with initial exposure evidence signal.\n\n"
        f"8. Missing Skills for Target Career:\n{missing_career_skills_str}\n\n"
        f"9. Recommendations for Resume Improvement:\n• Highlight measurable outcomes (metrics, % improvements, throughput).\n• Link GitHub repository URLs or live demo links for projects.\n• Take skill diagnostic assessments to verify claimed competencies to 80%+ mastery.\n\n"
        f"⚠️ Note: Resume detection reflects document claims and provides exposure calibration; it does not constitute independently verified mastery."
    )

    active_resume_id = resume_id or f"RES-{uuid.uuid4().hex[:8].upper()}"
    filesize = file_path.stat().st_size if file_path.exists() else 0
    now_ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")

    # Deactivate previous active resumes for this student
    db.query(ResumeDocument).filter(
        ResumeDocument.student_id == student_id,
        ResumeDocument.is_active == True,
    ).update({"is_active": False})

    resume_doc = ResumeDocument(
        resume_id=active_resume_id,
        student_id=student_id,
        original_filename=original_filename,
        file_path=str(file_path),
        file_size_bytes=filesize,
        upload_timestamp=now_ts,
        analysis_status="analyzed",
        analysis_version="v1.0",
        is_active=True,
        extracted_skills_json=json.dumps(detected_skills_list),
        extracted_projects_json=json.dumps([p.model_dump() for p in extracted_projects]),
        extracted_certificates_json=json.dumps([c.model_dump() for c in extracted_certificates]),
        extracted_experience_json=json.dumps([e.model_dump() for e in extracted_experience]),
        extracted_education_json=json.dumps([ed.model_dump() for ed in extracted_education]),
        analysis_summary=summary_text,
    )
    db.add(resume_doc)
    db.commit()

    return {
        "student_id": student_id,
        "resume_id": active_resume_id,
        "upload_timestamp": now_ts,
        "filename": original_filename,
        "total_skills_detected": len(detected_skills_list),
        "total_projects_detected": len(extracted_projects),
        "total_certificates_detected": len(extracted_certificates),
        "total_experience_detected": len(extracted_experience),
        "detected_skills": detected_skills_list,
        "unmatched_skills": unmatched_skills,
        "extracted_projects": [p.model_dump() for p in extracted_projects],
        "extracted_certificates": [c.model_dump() for c in extracted_certificates],
        "extracted_experience": [e.model_dump() for e in extracted_experience],
        "extracted_education": [ed.model_dump() for ed in extracted_education],
        "analysis_summary": summary_text,
        "disclaimer": "Resume evidence indicates demonstrated exposure, not verified mastery. Mastery requires assessment and project proof.",
    }

